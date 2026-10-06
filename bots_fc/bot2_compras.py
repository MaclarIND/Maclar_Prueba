"""Bot 2: comprador en la Web App de EA FC (mercado de transferencias).

Lee jugadores.xlsx (generado por el Bot 1) e intenta comprar cada jugador
respetando topes de precio. El login lo hacés VOS a mano.
"""
import re
import sys
import unicodedata
from datetime import datetime

from openpyxl import Workbook, load_workbook
from playwright.sync_api import sync_playwright

from comun import BASE, Parar, cargar_config, chequear_stop, asegurar_chromium, espera, log

URL_WEBAPP = "https://www.ea.com/ea-sports-fc/ultimate-team/web-app/"

# ======================================================================
# AJUSTES SI LA WEB APP CAMBIA (selectores y textos concentrados acá)
# ======================================================================
SEL = {
    # Monedas disponibles (arriba a la derecha)
    "monedas": ".view-navbar-currency-coins",
    # Sugerencias del autocompletado (cada una es un botón con nombre y rating)
    "sugerencias": ".playerResultsList button, .ut-player-search-control--results button",
    # Filas de resultados de la búsqueda
    "resultados": "ul.paginated-item-list > li, .ut-item-search-results li.listFocusable",
    # Botón "atrás" de la Web App
    "atras": "button.ut-navigation-button-control, .ut-navigation-bar-view button.btn-navigation",
}
TEXTO_MENU = ["Home", "Squads", "SBC", "Evolutions", "Transfers", "Store", "Club"]
PLACEHOLDER_NOMBRE = "Type Player Name"
TEXTO_TILE_BUSCAR = "Search the Transfer Market"
# Textos que indican captcha / bloqueo / error de EA -> el bot PARA y espera
REGEX_ALERTAS = re.compile(
    r"captcha|too many actions|too many requests|temporarily (banned|restricted)|"
    r"something went wrong|service unavailable|transfers? (are )?(currently )?disabled|"
    r"verify (that )?you are human", re.I)

# Encuentra el input "Max" de la sección "Buy Now Price:" (NO el de Bid Price) y su botón "+".
# Busca la etiqueta por texto y toma los inputs que vienen DESPUÉS de ella en el DOM.
JS_MARCAR_BUYNOW_MAX = r"""
() => {
  document.querySelectorAll('[data-bot]').forEach(e => e.removeAttribute('data-bot'));
  const FOLLOW = Node.DOCUMENT_POSITION_FOLLOWING;
  const etiquetas = [...document.querySelectorAll('body *')].filter(e =>
      e.children.length === 0 && e.offsetParent && /^buy now price/i.test(e.textContent.trim()));
  for (const et of etiquetas) {
    let nodo = et;
    let inputs = [];
    while (nodo && inputs.length < 2) {
      nodo = nodo.parentElement;
      if (!nodo) break;
      inputs = [...nodo.querySelectorAll('input')].filter(i => et.compareDocumentPosition(i) & FOLLOW);
    }
    if (inputs.length < 2) continue;
    const max = inputs[1];            // [0] = Min, [1] = Max
    max.setAttribute('data-bot', 'buynow-max');
    const botones = [...nodo.querySelectorAll('button')].filter(b =>
        (max.compareDocumentPosition(b) & FOLLOW) &&
        (b.textContent.trim() === '+' || /increment|plus/i.test(b.className)));
    if (botones[0]) botones[0].setAttribute('data-bot', 'buynow-plus');
    return true;
  }
  return false;
}
"""
# ======================================================================

def norm(s):
    s = unicodedata.normalize("NFD", s or "")
    return "".join(c for c in s if unicodedata.category(c) != "Mn").lower().strip()


def a_numero(txt):
    d = re.sub(r"[^\d]", "", txt or "")
    return int(d) if d else None


class Bot:
    def __init__(self, page, cfg):
        self.page = page
        self.cfg = cfg
        self.busquedas = 0
        self.gastado = 0
        self.resultados = []

    # ---------- utilidades ----------
    def esperar_busqueda(self):
        espera(self.cfg["espera_min_seg"], self.cfg["espera_max_seg"])

    def revisar_alertas(self):
        """Si EA muestra captcha/error: parar y esperar al humano. No se intenta saltear nada."""
        try:
            texto = self.page.inner_text("body", timeout=3000)
        except Exception:
            return
        m = REGEX_ALERTAS.search(texto)
        if m:
            log(f"!!! EA mostró una alerta: '{m.group(0)}'. El bot se PAUSA.")
            log("    Resolvelo vos en el navegador y apretá ENTER acá para continuar (o creá STOP.txt para cortar).")
            input()
            chequear_stop()

    def monedas(self):
        try:
            return a_numero(self.page.locator(SEL["monedas"]).first.inner_text(timeout=4000))
        except Exception:
            return None

    def click_texto(self, texto, exacto=True, timeout=8000):
        self.page.get_by_text(texto, exact=exacto).first.click(timeout=timeout)

    def click_boton(self, nombre, timeout=8000):
        self.page.get_by_role("button", name=re.compile(rf"^\s*{re.escape(nombre)}\s*$", re.I)).first.click(timeout=timeout)

    def campo_nombre(self):
        return self.page.get_by_placeholder(PLACEHOLDER_NOMBRE)

    # ---------- inicio ----------
    def esperar_login_y_home(self):
        log("Esperando que inicies sesión a mano (usuario, clave y código). No toques nada más...")
        while True:
            chequear_stop()
            try:
                lineas = {l.strip() for l in self.page.inner_text("body", timeout=2000).split("\n")}
                if all(t in lineas for t in TEXTO_MENU):
                    break
            except Exception:
                pass
            self.page.wait_for_timeout(1500)
        log("Home detectada. Arranco.")
        espera(2, 4)

    def ir_a_buscar(self):
        self.click_texto("Transfers")
        espera(1.5, 2.5)
        self.click_texto(TEXTO_TILE_BUSCAR, exacto=False)
        self.campo_nombre().wait_for(timeout=15000)

    def volver_a_filtros(self):
        for _ in range(3):
            if self.campo_nombre().count() and self.campo_nombre().first.is_visible():
                return
            self.page.locator(SEL["atras"]).first.click(timeout=5000)
            espera(0.8, 1.5)
        self.campo_nombre().wait_for(timeout=8000)

    # ---------- pasos por jugador ----------
    def escribir_nombre(self, nombre):
        campo = self.campo_nombre()
        campo.click()
        campo.fill("")
        campo.press_sequentially(nombre, delay=self.cfg["delay_tecla_ms"])
        self.page.wait_for_timeout(1500)

    def elegir_sugerencia(self, nombre, rating):
        """Elige la sugerencia cuyo texto contiene el nombre y (si hay) el rating."""
        sug = self.page.locator(SEL["sugerencias"])
        n = sug.count()
        candidatas = []
        for i in range(n):
            t = sug.nth(i).inner_text()
            if norm(nombre) in norm(t):
                candidatas.append((i, t))
        if rating:
            candidatas = [c for c in candidatas if re.search(rf"(?<!\d){rating}(?!\d)", c[1])]
        elif len(candidatas) > 1:
            log("  (sin rating en el Excel y hay varias sugerencias: tomo la primera)")
        if not candidatas:
            return False
        sug.nth(candidatas[0][0]).click()
        self.page.wait_for_timeout(800)
        return True

    def fijar_max_buynow(self, precio):
        if not self.page.evaluate(JS_MARCAR_BUYNOW_MAX):
            raise RuntimeError("No encuentro el campo Max de 'Buy Now Price' (ajustar JS_MARCAR_BUYNOW_MAX)")
        campo = self.page.locator('[data-bot="buynow-max"]')
        campo.click()
        campo.press("Control+A")
        campo.press("Backspace")
        campo.press_sequentially(str(precio), delay=60)
        campo.press("Tab")
        self.page.wait_for_timeout(400)

    def leer_max_buynow(self):
        self.page.evaluate(JS_MARCAR_BUYNOW_MAX)
        return a_numero(self.page.locator('[data-bot="buynow-max"]').input_value())

    def subir_max_con_mas(self):
        """Aprieta '+' del Max de Buy Now (EA respeta sus escalones). Devuelve el nuevo valor."""
        antes = self.leer_max_buynow() or 0
        self.page.locator('[data-bot="buynow-plus"]').click(timeout=5000)
        self.page.wait_for_timeout(300)
        despues = self.leer_max_buynow()
        if despues is None or despues <= antes:
            raise RuntimeError("El botón '+' no subió el precio")
        return despues

    def buscar(self):
        """Click en Search. Devuelve lista de (precio_buynow, indice_fila) o [] si 'No results'."""
        self.busquedas += 1
        self.click_boton("Search")
        self.page.wait_for_timeout(2500)
        self.revisar_alertas()
        filas = self.page.locator(SEL["resultados"])
        out = []
        for i in range(filas.count()):
            m = re.search(r"Buy Now(?: Price)?:?\s*([\d.,]+)", filas.nth(i).inner_text(), re.I)
            if m and a_numero(m.group(1)):
                out.append((a_numero(m.group(1)), i))
        return sorted(out)

    def comprar(self, fila_idx, precio):
        """Compra la fila indicada. Devuelve precio pagado (o precio listado en modo prueba)."""
        if self.cfg["modo_prueba"]:
            log(f"  [PRUEBA] Compraría por {precio} (no se aprieta Buy Now)")
            return precio, "Simulado (prueba)"
        antes = self.monedas()
        self.page.locator(SEL["resultados"]).nth(fila_idx).click()
        self.page.wait_for_timeout(800)
        self.click_boton("Buy Now")
        self.page.wait_for_timeout(800)
        self.click_boton("Ok")
        self.page.wait_for_timeout(2500)
        self.revisar_alertas()
        despues = self.monedas()
        pagado = antes - despues if antes is not None and despues is not None and antes > despues else precio
        return pagado, "Comprado"

    # ---------- flujo principal por jugador ----------
    def procesar(self, nombre, precio_xl, rating, equipo):
        tope = min(precio_xl * (1 + self.cfg["margen_max"]), self.cfg["precio_max_absoluto"])
        if precio_xl > tope:
            return None, "Tope superado"
        self.volver_a_filtros()
        self.click_boton("Reset")
        espera(0.5, 1)
        self.escribir_nombre(nombre)
        if not self.elegir_sugerencia(nombre, rating):
            log("  No encontrado en autocompletado")
            return None, "No encontrado"
        precio_actual = precio_xl
        self.fijar_max_buynow(precio_actual)
        while True:
            if self.busquedas >= self.cfg["max_busquedas_sesion"]:
                raise Parar("Límite de búsquedas por sesión alcanzado")
            if self.busquedas and self.busquedas % self.cfg["pausa_larga_cada"] == 0:
                log(f"Pausa larga de {self.cfg['pausa_larga_seg']} s (anti-bloqueo)")
                espera(self.cfg["pausa_larga_seg"])
            self.esperar_busqueda()
            resultados = self.buscar()
            if resultados:
                menor, idx = resultados[0]
                if menor > tope:
                    return None, "Tope superado"
                if self.gastado + menor > self.cfg["presupuesto_total"]:
                    log("  Superaría el presupuesto_total")
                    return None, "Tope superado"
                monedas = self.monedas()
                if monedas is not None and monedas < menor:
                    raise Parar(f"Monedas insuficientes ({monedas} < {menor})")
                pagado, estado = self.comprar(idx, menor)
                self.gastado += pagado
                return pagado, estado
            # "No results": volver y subir el Max con '+'
            log(f"  Sin resultados con Max {precio_actual}")
            self.volver_a_filtros()
            precio_actual = self.subir_max_con_mas()
            if precio_actual > tope:
                log(f"  Tope alcanzado ({precio_actual} > {tope:.0f})")
                return None, "Tope superado"

    def guardar(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Compras"
        ws.append(["Nombre", "Precio Excel", "Precio pagado", "Estado", "Hora"])
        for r in self.resultados:
            ws.append(r)
        wb.save(BASE / "compras_resultado.xlsx")


def leer_excel(ruta):
    ws = load_workbook(ruta, read_only=True).active
    filas = []
    for nombre, precio, rating, equipo in ws.iter_rows(min_row=2, max_col=4, values_only=True):
        if nombre and precio:
            filas.append((str(nombre), int(precio), int(rating) if rating else None, equipo))
    return filas


def main():
    cfg = cargar_config()
    ruta = BASE / cfg["excel_entrada"]
    if not ruta.exists():
        raise SystemExit(f"No encuentro {ruta}. Corré primero bot1_futgg.exe")
    jugadores = leer_excel(ruta)
    log(f"{len(jugadores)} jugadores en el Excel. modo_prueba={cfg['modo_prueba']}")
    asegurar_chromium()
    with sync_playwright() as p:
        ctx = p.chromium.launch_persistent_context(
            str(BASE / "perfil_chrome"), headless=False, viewport=None, args=["--start-maximized"])
        page = ctx.pages[0] if ctx.pages else ctx.new_page()
        bot = Bot(page, cfg)
        try:
            page.goto(URL_WEBAPP)
            bot.esperar_login_y_home()
            bot.ir_a_buscar()
            for nombre, precio, rating, equipo in jugadores:
                chequear_stop()
                log(f"== {nombre} ({rating}, {equipo}) precio Excel {precio}")
                try:
                    pagado, estado = bot.procesar(nombre, precio, rating, equipo)
                except Parar:
                    raise
                except Exception as e:
                    log(f"  ERROR: {e}")
                    pagado, estado = None, "Error"
                    bot.revisar_alertas()  # si fue por un bloqueo, esperar al humano
                log(f"  -> {estado}" + (f" ({pagado})" if pagado else ""))
                bot.resultados.append([nombre, precio, pagado, estado, datetime.now().strftime("%H:%M:%S")])
                bot.guardar()
        except Parar as e:
            log(f"PARADA: {e}")
        finally:
            bot.guardar()
            log(f"Resultados en {BASE / 'compras_resultado.xlsx'} | gastado: {bot.gastado}")
            ctx.close()


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        pass
    input("\nEnter para cerrar...") if getattr(sys, "frozen", False) else None
